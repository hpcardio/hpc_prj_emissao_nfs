from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from nfs_fortaleza.maida_portal import (
    DownloadedMaidaDocument,
    MaidaCompetency,
    MaidaDocument,
)
from nfs_fortaleza.maida_xlsx import MaidaSpreadsheetError, parse_maida_xlsx


def _download(path: Path) -> DownloadedMaidaDocument:
    return DownloadedMaidaDocument(
        document=MaidaDocument(
            competency=MaidaCompetency(2026, 1),
            source="resumo",
            document_id="doc-1",
            file_name="demonstrativo.xlsx",
            remote_path="/arquivo/demonstrativo.xlsx",
        ),
        path=path,
    )


def test_parse_maida_xlsx_filters_and_normalizes_glosas(tmp_path: Path) -> None:
    path = tmp_path / "demonstrativo.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Demonstrativo ISSEC"])
    sheet.append(
        [
            "NRO PROTOCOLO",
            "NÚMERO GUIA",
            "PACIENTE",
            "COD. PROCEDIMENTO",
            "DESCRIÇÃO PROCEDIMENTO",
            "QTD EXECUTADA",
            "VALOR APRESENTADO",
            "VALOR PAGO",
            "VALOR GLOSA",
            "CÓDIGO GLOSA",
            "DATA REALIZAÇÃO",
        ]
    )
    sheet.append(
        [
            "5583152",
            "521882",
            "PACIENTE TESTE",
            "30912199",
            "PROCEDIMENTO TESTE",
            2,
            "1.200,50",
            "900,25",
            "300,25",
            "1714 - VALOR SUPERIOR",
            "05/01/2026",
        ]
    )
    sheet.append(["sem-glosa", None, None, None, None, 1, 10, 10, 0, 1714])
    sheet.append(["sem-codigo", None, None, None, None, 1, 10, 0, 10, None])
    workbook.save(path)

    records = list(parse_maida_xlsx(_download(path)))

    assert len(records) == 1
    assert records[0]["numero_protocolo"] == "5583152"
    assert records[0]["codigo_glosa"] == "1714"
    assert records[0]["codigo_servico"] == "30912199"
    assert records[0]["valor_glosa"] == Decimal("300.25")
    assert records[0]["valor_processado"] == Decimal("1200.50")
    assert records[0]["valor_liberado"] == Decimal("900.25")
    assert records[0]["data_realizacao"] == date(2026, 1, 5)
    assert records[0]["referencia"] == date(2026, 1, 1)
    assert records[0]["id_registro"]


def test_parse_maida_xlsx_rejects_unknown_layout(tmp_path: Path) -> None:
    path = tmp_path / "invalido.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["COLUNA A", "COLUNA B"])
    workbook.save(path)

    with pytest.raises(MaidaSpreadsheetError, match="Valor Glosa"):
        list(parse_maida_xlsx(_download(path)))
